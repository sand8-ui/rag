"""Generate StayWise RAG corpus in DOCX. Topics do not overlap markdown files."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = Path(__file__).parent / "docx"
OUT.mkdir(parents=True, exist_ok=True)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_paras(doc: Document, paragraphs: list[str]) -> None:
    for text in paragraphs:
        doc.add_paragraph(text)


def save(doc: Document, name: str) -> None:
    path = OUT / name
    doc.save(path)
    print(f"wrote {path}")


def doc_change_dates() -> None:
    doc = Document()
    title = doc.add_paragraph("StayWise 订单改期作业手册")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_paras(
        doc,
        [
            "文档编号：SW-POL-CHANGE-2026。本文件只讲改期，不讲取消（取消见 markdown 取消政策）。",
            "演示环境改期入口：登录后「我的订单」→「改期」。接口为 PATCH /orders/:id，请求体仅含 checkIn、checkOut。",
        ],
    )
    add_heading(doc, "1. 允许改什么", 1)
    add_paras(
        doc,
        [
            "允许：入住日期、离店日期。离店必须严格晚于入住；入住不得早于改期操作当天（上海时区日历日）。",
            "不允许：更换房型、更换酒店、修改人数、修改价格、修改入住人姓名。这些需求要删旧单再订，或转人工。",
            "改期不重新计算房价。快照里的每晚价格保持下单时写入的数字，即使日历房已变价。",
        ],
    )
    add_heading(doc, "2. 次数与费用", 1)
    add_paras(
        doc,
        [
            "限时免费取消产品：入住前 1 个自然日 23:59 前，改期免费，不限次数，但两次改期间隔建议不少于 10 分钟，避免误触。",
            "入住日当天改期：视为占用当天库存，可能按「取消再预订」处理；演示系统仍只改日期字段，客服应提示「现场可能按新日期重新确认」。",
            "不可取消产品：原则上不能改期。用户坚持则转人工，可能收取一晚差价或拒绝。",
        ],
    )
    add_heading(doc, "3. 与搜索条件的关系", 1)
    add_paras(
        doc,
        [
            "用户从搜索页带入的日期只是预填。改期后不会自动回写搜索页。",
            "改完日期不校验新日期是否与其他客人撞房。演示无库存锁。",
        ],
    )
    add_heading(doc, "4. 失败场景", 1)
    add_paras(
        doc,
        [
            "400：日期格式不是 YYYY-MM-DD、离店不晚于入住、入住早于今天。",
            "401：未登录或 access token 过期。",
            "404：订单不属于当前用户，或已被删除。",
        ],
    )
    save(doc, "01-order-date-change.docx")


def doc_invoice() -> None:
    doc = Document()
    doc.add_paragraph("StayWise 发票与报销指引").alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_paras(
        doc,
        [
            "文档编号：SW-POL-INV-2026。本文件不讨论取消扣费比例。",
            "演示环境没有自动开票接口。客服只能收集信息并说明「上线后的流程」。",
        ],
    )
    add_heading(doc, "1. 谁来开票", 1)
    add_paras(
        doc,
        [
            "StayWise 预付：由平台开具「代订住宿服务费」或「旅游服务费」类发票，不是酒店住宿专票。企业报销前请用户向财务确认能否入账。",
            "到店现付房费：由酒店开具住宿发票。平台不应重复开同一笔房费。",
            "押金、迷你吧、SPA：一律酒店现场开具，平台不能代开。",
        ],
    )
    add_heading(doc, "2. 需要用户提供的字段", 1)
    add_paras(
        doc,
        [
            "抬头类型：企业或个人。企业需完整名称与纳税人识别号。",
            "邮箱：用于接收电子发票 PDF。",
            "订单号、入住人、入住离店日期、税前金额。",
            "是否需要「住宿天数」备注。部分单位要求备注房型。",
        ],
    )
    add_heading(doc, "3. 时限", 1)
    add_paras(
        doc,
        [
            "建议离店后 30 日内申请。超过 90 日需财务人工审批。",
            "电子票一般 3 个工作日发送。纸质票仅对公客户，快递到付。",
            "换开：电子票红冲后重开，每月限 1 次。",
        ],
    )
    add_heading(doc, "4. 积分与优惠", 1)
    add_paras(
        doc,
        [
            "积分抵扣部分不开票。发票金额 = 用户实付现金。",
            "不可开「含积分的全额房费票」。",
        ],
    )
    save(doc, "02-invoice-reimbursement.docx")


def doc_payment() -> None:
    doc = Document()
    doc.add_paragraph("StayWise 支付、担保与退款路径").alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_paras(
        doc,
        [
            "文档编号：SW-POL-PAY-2026。演示尚未接入收银台，本文件描述目标产品口径，避免客服说「已经扣款」或「肯定没扣款」。",
        ],
    )
    add_heading(doc, "1. 当前演示", 1)
    add_paras(
        doc,
        [
            "下单不经过支付网关，订单直接进入数据库。用户问「我的卡被刷了吗」：回答演示环境不会发生真实扣款。",
            "前台仍可能收押金预授权，那是酒店行为，与 StayWise 演示支付无关。",
        ],
    )
    add_heading(doc, "2. 上线后的三种模式", 1)
    add_paras(
        doc,
        [
            "预付：下单时支付全部房费到 StayWise 监管账户，离店后与酒店结算。",
            "到店付：下单仅锁意向，入住时向酒店付房费。取消规则仍按产品标签。",
            "信用卡担保：下单时预授权 1 晚，No-Show 时请款。",
        ],
    )
    add_heading(doc, "3. 退款路径", 1)
    add_paras(
        doc,
        [
            "原路退回。信用卡 1–15 工作日，借记卡与钱包通常更快。",
            "演示删除订单不会触发退款流水。",
            "争议单保留：支付单号、银行短信、订单截图。",
        ],
    )
    add_heading(doc, "4. 外币", 1)
    add_paras(
        doc,
        [
            "演示仅人民币标价。外卡在酒店刷押金时的汇率以收单行账单为准。",
        ],
    )
    save(doc, "03-payment-guarantee.docx")


def doc_pets() -> None:
    doc = Document()
    doc.add_paragraph("StayWise 宠物与过敏协办手册").alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_paras(
        doc,
        [
            "文档编号：SW-POL-PET-2026。服务犬规则见 markdown《无障碍与特殊需求》。本文件只谈伴侣宠物。",
        ],
    )
    add_heading(doc, "1. 四家店宠物政策（演示口径）", 1)
    add_paras(
        doc,
        [
            "上海外滩华尔道夫：原则上不接受宠物，服务犬除外。擅自带入可能被拒住且不予退款。",
            "杭州西溪悦榕庄：小型犬（体重建议 8kg 以下）可能接受，需至少提前 48 小时申请，清洁费约 300–600 元/晚，现场付。猫视供应。",
            "成都太古里博舍：不接受宠物。商圈室内商场亦常禁宠。",
            "北京王府井文华东方：不接受宠物。城区豪华酒店对毛发清洁要求高。",
        ],
    )
    add_heading(doc, "2. 申请要写清", 1)
    add_paras(
        doc,
        [
            "物种、体重、是否具备疫苗本、是否驱虫、是否会单独留在房间。",
            "不可将宠物单独留房超过 2 小时。禁止宠物上公共床品；酒店可提供宠物垫。",
            "损坏与扰民由客人承担。",
        ],
    )
    add_heading(doc, "3. 过敏客人", 1)
    add_paras(
        doc,
        [
            "即使某店禁宠，历史房间仍可能有毛发残留。严重过敏请在工单里写「需无宠楼层」，不保证绝对无毛。",
        ],
    )
    save(doc, "04-pets-and-allergies.docx")


def doc_dining() -> None:
    doc = Document()
    doc.add_paragraph("StayWise 早餐、餐饮与迷你吧说明").alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_heading(doc, "1. 早餐", 1)
    add_paras(
        doc,
        [
            "演示房价默认不含早餐。用户问「多少钱一份」：上海约 180–280，杭州度假店约 200–320，成都约 120–198，北京约 220–350，均为约数，以下单或前台为准。",
            "早餐时段常见 6:30–10:00，周末或节假日可能延至 10:30。过时不补、不打包，除非酒店明确允许。",
            "儿童早餐：3 岁以下可能免费，3–11 岁半价或固定价，12 岁以上成人价。以酒店当天牌价为准。",
        ],
    )
    add_heading(doc, "2. 餐厅预约", 1)
    add_paras(
        doc,
        [
            "江景西餐厅、湿地餐厅周五周六需预约。StayWise 可代为转达，成功与否回传给用户。",
            "不代收订金。放鸽子可能影响后续预约。",
        ],
    )
    add_heading(doc, "3. 迷你吧与客房送餐", 1)
    add_paras(
        doc,
        [
            "迷你吧按酒水单计费，从押金扣除。争议时看小条与冰箱传感器记录。",
            "客房送餐通常 6:00–23:00，加收服务费 10–15%。",
            "禁止外带火锅、烧烤进房加热。触发报警的处置费由客人承担。",
        ],
    )
    add_heading(doc, "4. 宗教与饮食禁忌", 1)
    add_paras(
        doc,
        [
            "可备注素食、清真、花生过敏。酒店尽量协助，不保证独立厨房。",
        ],
    )
    save(doc, "05-dining-and-minibar.docx")


def doc_safety() -> None:
    doc = Document()
    doc.add_paragraph("StayWise 安全、贵重物品与保险须知").alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_heading(doc, "1. 贵重物品", 1)
    add_paras(
        doc,
        [
            "房间保险箱仅供便利，丢失按酒店住宿合同处理。大额现金、护照建议使用前台保管。",
            "StayWise 不是保险公司，不赔付房间内失窃。可提示用户联系信用卡盗刷保障或旅游险。",
        ],
    )
    add_heading(doc, "2. 消防与逃生", 1)
    add_paras(
        doc,
        [
            "入住时请用户看房门内侧逃生图。禁止封堵烟感。",
            "历史建筑（上海外滩店）走廊曲折，更要确认最近楼梯。",
        ],
    )
    add_heading(doc, "3. 意外", 1)
    add_paras(
        doc,
        [
            "滑倒、泳池、交通：先保障人身安全，再报酒店安保与客服。不要在聊天里做医疗诊断。",
            "用户索要「平台赔偿医疗费」必须转人工法务，禁止 RAG 承诺金额。",
        ],
    )
    add_heading(doc, "4. 监控与隐私", 1)
    add_paras(
        doc,
        [
            "公共区域可能有监控。客房、卫生间无监控。",
            "不要向第三方泄露用户证件号与订单手机号。",
        ],
    )
    save(doc, "06-safety-and-valuables.docx")


def doc_business() -> None:
    doc = Document()
    doc.add_paragraph("StayWise 商务出行、会议与企业协议").alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_paras(
        doc,
        [
            "文档编号：SW-POL-BIZ-2026。演示无企业后台。遇「我们公司有协议价」一律转人工，不要用散客价冒充协议价。",
        ],
    )
    add_heading(doc, "1. 会议与宴会", 1)
    add_paras(
        doc,
        [
            "四家店均具备或可协调小型会议室，具体桌型与价格现场报价。",
            "西溪更适合团建户外；王府井、外滩适合政务接待；太古里适合小型发布。",
            "需要投影、同传、茶歇：列表需求后由酒店会议销售回复，StayWise 只做传递。",
        ],
    )
    add_heading(doc, "2. 连住与团队房", 1)
    add_paras(
        doc,
        [
            "5 间及以上视为团队，可能要求预付或签订 BEO。演示系统按散客单处理，需拆单或人工建团。",
            "团队取消窗口通常比散客更早，常见入住前 14 天。",
        ],
    )
    add_heading(doc, "3. 出行单与对公结算", 1)
    add_paras(
        doc,
        [
            "月结账户不在演示范围。用户发「我们是月结客户」时收集公司全称与历史协议编号，转商务组。",
        ],
    )
    save(doc, "07-business-meetings.docx")


def doc_housekeeping() -> None:
    doc = Document()
    doc.add_paragraph("StayWise 客房清洁、续住与物品遗失").alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_heading(doc, "1. 清洁时段", 1)
    add_paras(
        doc,
        [
            "日常清洁大约 9:00–15:00。挂「请勿打扰」则跳过当天。连续多日 DND 时，酒店可能按消防要求敲门确认安全。",
            "环保选项：连续住可申请隔日换布草。",
        ],
    )
    add_heading(doc, "2. 续住", 1)
    add_paras(
        doc,
        [
            "演示改期可以把离店日延后，不保证仍是原房号。满房时酒店可能安排换房。",
            "续住价格沿用订单快照单价，不自动改成当天门市价。",
        ],
    )
    add_heading(doc, "3. 遗失物品", 1)
    add_paras(
        doc,
        [
            "离店后发现遗失：提供房号或订单号、物品特征、大概位置。酒店保管期常见 30–90 天。",
            "寄回运费到付。液体、食品、低值牙刷类通常不寄。",
            "StayWise 不能强制酒店承认「一定找到」。",
        ],
    )
    add_heading(doc, "4. 损坏", 1)
    add_paras(
        doc,
        [
            "床单污染、地毯烧痕、设备损坏按酒店报价从押金扣。用户质疑报价时要照片与价目表，不要帮用户骂酒店。",
        ],
    )
    save(doc, "08-housekeeping-lost-found.docx")


def doc_transfer_parking() -> None:
    doc = Document()
    doc.add_paragraph("StayWise 接送机、停车与充电").alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_paras(
        doc,
        [
            "文档编号：SW-POL-TRANS-2026。演示订单不含免费接机。本文件回答「有没有车接」「停车多少钱」「能充电吗」。",
        ],
    )
    add_heading(doc, "1. 接机接站", 1)
    add_paras(
        doc,
        [
            "四家店均可代订收费接送，价格随车型与时段浮动，必须向酒店礼宾确认。StayWise 不代收款、不派自有司机。",
            "上海：虹桥通常比浦东便宜。高峰加价。外滩下车点可能临时管制。",
            "杭州：西溪内部可能要在门外换乘接驳，司机不能开到别墅门口。",
            "成都：天府机场显著远于双流，报价前先问清机场。",
            "北京：活动日长安街管制，接机只能送到附近路口。",
            "航班延误：用户需自己改预约。司机空等超时费与平台无关。",
        ],
    )
    add_heading(doc, "2. 停车", 1)
    add_paras(
        doc,
        [
            "外滩：代客泊车为主，过夜费约 150–200 元量级，以现场为准。车位极度紧张。",
            "西溪：相对好停，仍可能下雨后车位泥泞，提醒用户。",
            "太古里：商圈车库贵且满，建议地铁。",
            "王府井：地下库有，大型活动可能临时封闭。",
            "新能源：不要承诺一定有空闲充电桩。有桩也可能要排队或临时收费。",
        ],
    )
    add_heading(doc, "3. 客服禁语", 1)
    add_paras(
        doc,
        [
            "不说「司机已经出发」。不说「停车免费」。不替用户叫以个人微信收款的黑车。",
        ],
    )
    save(doc, "09-airport-transfer-parking.docx")


def doc_upgrades() -> None:
    doc = Document()
    doc.add_paragraph("StayWise 升房、特殊布置与迎宾").alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_heading(doc, "1. 免费升房", 1)
    add_paras(
        doc,
        [
            "视当天房态，前台决定。StayWise 不能在演示里锁定升级。",
            "会员等级「行家/领航」仅增加申请优先级，不是权利。",
            "用户已订最高档（如紫禁城景观套房、湿地别墅）时，没有更高房型可升。",
        ],
    )
    add_heading(doc, "2. 付费升级", 1)
    add_paras(
        doc,
        [
            "差价按酒店门市或前台报价，现场付。平台订单价格不自动改。",
            "升级后原取消规则通常仍跟原订单产品走。",
        ],
    )
    add_heading(doc, "3. 生日、求婚、纪念日", 1)
    add_paras(
        doc,
        [
            "可转达鲜花、蛋糕、摆花。费用酒店收。不保证保密成功（别把惊喜发到入住人邮箱抄送）。",
            "外滩江景求婚：人多、安保严，户外大型布置可能被拦，建议室内。",
            "禁止承诺「一定给最好的角度」。",
        ],
    )
    add_heading(doc, "4. 枕头被褥与气味", 1)
    add_paras(
        doc,
        [
            "额外枕头、被子、冷气被：联系管家即可，一般免费。",
            "无羽绒、乳胶替换：尽量满足，数量有限。",
        ],
    )
    save(doc, "10-upgrade-special-requests.docx")


def doc_fraud() -> None:
    doc = Document()
    doc.add_paragraph("StayWise 防诈骗与假客服识别").alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_paras(
        doc,
        [
            "文档编号：SW-SOP-FRAUD-2026。用户说「有人打电话让我补差价」时用本篇。",
        ],
    )
    add_heading(doc, "1. 真假辨别", 1)
    add_paras(
        doc,
        [
            "StayWise 不会要用户把房费转到个人微信、支付宝、比特币。",
            "不会要密码、短信验证码、银行卡有效期和 CVV。演示甚至没有短信验证码。",
            "不会用「公安 / 法院 / 酒店财务」施压立即转账。",
            "链接如果不是用户自己打开的 StayWise 页面，先怀疑。",
        ],
    )
    add_heading(doc, "2. 常见剧本", 1)
    add_paras(
        doc,
        [
            "「系统重复扣款，我帮你退，先转 1 元验证。」骗子。",
            "「你的订单异常，点这个短链改密。」钓鱼。",
            "「酒店改名收款账户。」让用户到前台看官方收款码，不要信电话里的卡号。",
        ],
    )
    add_heading(doc, "3. 已经转了钱", 1)
    add_paras(
        doc,
        [
            "立即联系银行/支付机构。保留通话记录。平台协助报警材料，但不垫付。",
            "AI 禁止帮用户「试试再转一笔追回」。",
        ],
    )
    save(doc, "11-fraud-and-phishing.docx")


def doc_visa_letter() -> None:
    doc = Document()
    doc.add_paragraph("StayWise 签证确认函与英文证明").alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_paras(
        doc,
        [
            "文档编号：SW-POL-VISA-2026。留学生、外籍旅客、出差办签常要「酒店预订单」。",
        ],
    )
    add_heading(doc, "1. 演示能给什么", 1)
    add_paras(
        doc,
        [
            "用户可自行截图「我的订单」。这不是酒店盖章确认函，使馆不一定收。",
            "StayWise 演示不能生成带公章 PDF，不能出具邀请函，不能担保入境。",
        ],
    )
    add_heading(doc, "2. 上线后规划", 1)
    add_paras(
        doc,
        [
            "英文确认函字段：客人姓名（护照）、酒店英文名与地址、入住离店、房型、订单号、已预付/到店付。",
            "不可取消订单更易被使馆采信，但客服不要为了签证怂恿用户买不可退产品而不说明风险。",
            "改期后旧函作废，需重开。",
        ],
    )
    add_heading(doc, "3. 姓名必须对护照", 1)
    add_paras(
        doc,
        [
            "拼音错误会导致拒签或无法入住。先改名再出函。演示无改名接口，见 markdown 代订与姓名文档。",
        ],
    )
    save(doc, "12-visa-confirmation.docx")


if __name__ == "__main__":
    doc_change_dates()
    doc_invoice()
    doc_payment()
    doc_pets()
    doc_dining()
    doc_safety()
    doc_business()
    doc_housekeeping()
    doc_transfer_parking()
    doc_upgrades()
    doc_fraud()
    doc_visa_letter()
    print("done")
